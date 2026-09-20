# -*- coding: utf-8 -*-
"""Cury: DRE trimestral (receita, lucro bruto, despesas, lucro antes do resultado financeiro, lucro líquido) e balanço (PL total,
PL da controladora, ativo total) da planilha Fundamentos do RI (fontes/verificacao/cury_planilha_fundamentos.xlsx), R$ mil.
Saída: _cury_dre.json. Usado no slide 'Retorno por vertical' como benchmark do MCMV."""
import openpyxl, json, io, datetime, os
here = os.path.dirname(os.path.abspath(__file__))
wb = openpyxl.load_workbook(os.path.join(here, "fontes", "verificacao", "cury_planilha_fundamentos.xlsx"), read_only=True, data_only=True)
ws = wb["Demonstrações do Resultado"]; rows = list(ws.iter_rows(values_only=True, max_row=45)); hdr = rows[1]
col = {str(h): i for i, h in enumerate(hdr) if h is not None and "T" in str(h) and len(str(h)) == 4}
out = {q: {"rec": rows[3][i], "lb": rows[5][i], "desp": rows[14][i], "lop": rows[15][i], "ll_ctrl": rows[29][i], "ll": rows[25][i]} for q, i in col.items()}
def bs(sheet, labels):
    ws = wb[sheet]; rows = list(ws.iter_rows(values_only=True, max_row=60)); hdr = rows[3]; res = {}
    for r in rows:
        lab = " ".join(str(c) for c in r[:3] if c)
        for key, pat in labels.items():
            if lab.upper().startswith(pat):
                for i, h in enumerate(hdr):
                    if isinstance(h, (datetime.datetime, datetime.date)) and isinstance(r[i], (int, float)):
                        res.setdefault(f"{(h.month - 1) // 3 + 1}T{h.year % 100:02d}", {})[key] = r[i]
    return res
P = bs("Balanço Patrimonial Passivo", {"pl_total": "PATRIMÔNIO LÍQUIDO TOTAL", "pl_ctrl": "PATRIMÔNIO LÍQUIDO DA CONTROLADORA"}); A = bs("Balanço Patrimonial Ativo", {"ativo": "TOTAL DO ATIVO"})
for q in out: out[q].update(P.get(q, {})); out[q].update(A.get(q, {}))
json.dump({"_meta": {"fonte": "Cury RI, planilha Fundamentos (DRE trimestral e balanço), R$ mil; lop = lucro antes do resultado financeiro; pl_total inclui não controladores"}, "serie": out}, io.open(os.path.join(here, "_cury_dre.json"), "w", encoding="utf-8"), ensure_ascii=False)
print("ok", len(out), "trimestres; sem balanço:", [q for q in out if "pl_total" not in out[q]])
